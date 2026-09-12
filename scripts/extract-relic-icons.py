"""Extract relic icons from the Slayer Legend master document.

Usage: python scripts/extract-relic-icons.py [path-to-xlsx]

Relic data itself comes from the wiki (fetch-wiki-equipment.py), but the
wiki's CDN does not name its relic art per relic, while the master document
anchors each icon to its own row. So the icons come from here, and
fetch-wiki-equipment.py matches them by relic id.

Run this before fetch-wiki-equipment.py.
"""

import re
import shutil
import sys
from pathlib import Path

import openpyxl

DEFAULT_SOURCE = Path.home() / "Downloads" / "Slayer Legend Master Document.xlsx"
ROOT = Path(__file__).resolve().parent.parent
RELIC_ICONS = ROOT / "public" / "relics"

ICON_COLUMN = 1  # zero-based: column B


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    if not source.exists():
        print(f"Cannot find {source}", file=sys.stderr)
        return 1

    sheet = openpyxl.load_workbook(source, data_only=True)["RELICS"]

    icons = {}
    for image in sheet._images:
        anchor = getattr(image.anchor, "_from", None)
        if anchor is not None and anchor.col == ICON_COLUMN:
            icons[anchor.row + 1] = image._data()

    if RELIC_ICONS.exists():
        shutil.rmtree(RELIC_ICONS)
    RELIC_ICONS.mkdir(parents=True)

    written = 0
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=3).value
        if not name or row not in icons:
            continue
        relic_id = int(sheet.cell(row=row, column=1).value)
        (RELIC_ICONS / f"{relic_id:02d}-{slug(str(name))}.png").write_bytes(icons[row])
        written += 1

    print(f"Wrote {written} relic icons to {RELIC_ICONS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
