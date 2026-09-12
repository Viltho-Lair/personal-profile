"""Extract the SKILLS sheet of the Slayer Legend master document into JSON.

Usage: python scripts/extract-skills.py [path-to-xlsx]

The sheet lays skills out as a grid: four columns (Fire, Water, Wind, Earth)
per row, running from the lowest grade to the highest. Row order is preserved.
Columns that are spreadsheet calculator scratch space rather than game data
(ENTER LEVEL, CURRENT STAT, DPS) are left out; they are derived from a level
the reader types in, and CURRENT STAT is just base + level * upgrade.
"""

import hashlib
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl

DEFAULT_SOURCE = Path.home() / "Downloads" / "Slayer Legend Master Document.xlsx"
ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "src" / "data" / "skills.json"
SKILL_ICONS = ROOT / "public" / "skills"
ELEMENT_ICONS = ROOT / "public" / "elements"

# Zero-based anchor columns of the two image sets embedded in the sheet.
SKILL_ICON_COLUMN = 1
ELEMENT_ICON_COLUMN = 3

ELEMENTS = ["Fire", "Water", "Wind", "Earth"]
GRADES = ["Common", "Great", "Rare", "Epic", "Legendary", "Mythic", "Immortal"]

# 1-based column numbers in the sheet.
COL = {
    "id": 1,
    "name": 3,
    "element": 5,
    "basic": 7,
    "specific": 8,
    "grade": 9,
    "maxLevel": 10,
    "openType": 11,
    "skillType": 12,
    "activeType": 13,
    "passiveType": 14,
    "mpCost": 15,
    "baseValue": 16,
    "upgradeValue": 17,
    "cooldown": 19,
    "range": 20,
    "duration": 21,
    "atkDistance": 22,
}


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def extract_images(sheet):
    """Return {sheet_row: image_bytes} for each anchored image column."""
    skills, elements = {}, {}
    for image in sheet._images:
        anchor = getattr(image.anchor, "_from", None)
        if anchor is None:
            continue
        row = anchor.row + 1  # openpyxl anchors are zero-based
        data = image._data()
        if anchor.col == SKILL_ICON_COLUMN:
            skills[row] = data
        elif anchor.col == ELEMENT_ICON_COLUMN:
            elements[row] = data
    return skills, elements


def number(value):
    """Whole floats become ints; everything else keeps its precision."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def text(value):
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    if not source.exists():
        print(f"Cannot find {source}", file=sys.stderr)
        return 1

    sheet = openpyxl.load_workbook(source, data_only=True)["SKILLS"]
    skill_images, element_images = extract_images(sheet)

    for folder in (SKILL_ICONS, ELEMENT_ICONS):
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)

    # The same four element icons repeat down the sheet, so write each once.
    element_icon_names: dict[str, str] = {}
    skills = []

    for row in range(2, sheet.max_row + 1):
        def cell(key):
            return sheet.cell(row=row, column=COL[key]).value

        name = text(cell("name"))
        if name is None:
            continue

        element = text(cell("element"))
        if element == "-":
            element = None

        skill_id = number(cell("id"))
        name_slug = slug(name)

        icon = None
        if row in skill_images:
            filename = f"{skill_id:02d}-{name_slug}.png"
            (SKILL_ICONS / filename).write_bytes(skill_images[row])
            icon = f"/skills/{filename}"

        element_icon = None
        if element and row in element_images:
            data = element_images[row]
            digest = hashlib.md5(data).hexdigest()
            filename = element_icon_names.get(digest)
            if filename is None:
                filename = f"{slug(element)}.png"
                (ELEMENT_ICONS / filename).write_bytes(data)
                element_icon_names[digest] = filename
            element_icon = f"/elements/{filename}"

        skills.append(
            {
                "id": skill_id,
                "row": (skill_id + 3) // 4,
                "name": name,
                "element": element,
                "icon": icon,
                "elementIcon": element_icon,
                "grade": text(cell("grade")),
                "maxLevel": number(cell("maxLevel")),
                "mpCost": number(cell("mpCost")),
                "baseValue": number(cell("baseValue")),
                "upgradeValue": number(cell("upgradeValue")),
                "cooldown": number(cell("cooldown")),
                "range": number(cell("range")),
                "duration": number(cell("duration")),
                "atkDistance": number(cell("atkDistance")),
                "types": {
                    "open": number(cell("openType")),
                    "skill": number(cell("skillType")),
                    "active": number(cell("activeType")),
                    "passive": number(cell("passiveType")),
                },
                "description": {
                    "basic": text(cell("basic")),
                    "specific": text(cell("specific")),
                },
            }
        )

    unknown_grades = sorted({s["grade"] for s in skills} - set(GRADES))
    if unknown_grades:
        print(f"Warning: grades not in the known order: {unknown_grades}", file=sys.stderr)

    data = {
        "source": {
            "file": source.name,
            "sheet": "SKILLS",
            "extractedOn": date.today().isoformat(),
        },
        "grades": GRADES,
        "elements": ELEMENTS,
        "skills": skills,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    icons = sum(1 for s in skills if s["icon"])
    print(f"Wrote {len(skills)} skills to {OUTPUT}")
    print(f"Wrote {icons} skill icons and {len(element_icon_names)} element icons")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
