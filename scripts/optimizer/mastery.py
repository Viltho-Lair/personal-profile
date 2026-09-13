"""SKILL MASTERY -> pages of nodes and the links drawn between them.

Each page is a 20-column block headed "PAGE n" on row 4. A node is a 5x4
block of cells:

    row r-1   icon, type label (ATK, EXP, SKILL, ...)
    row r     COST formula, "REWARD:" or a "REQUIRES:" formula   <- anchor
    row r+1   "BONUS :" and its value, or the reward / requirement text
    row r+2   the player's input: a level (validated 0..max) or a checkbox

Links are grey-filled cells on the page's black background. Grey cells that
touch each other form one connector; the nodes it touches are linked.
"""

import hashlib
import re

from openpyxl.utils import get_column_letter, range_boundaries

from optimizer.workbook import MissingHeader, images_by_cell, text

PAGE_WIDTH = 20
NODE_WIDTH, NODE_HEIGHT = 5, 4
LINK_FILL = "FF666666"
FIRST_ROW = 5

COST_LEVELS = re.compile(r'"COST:\s*"\s*&\s*(\d+(?:\.\d+)?)\s*-\s*[A-Z]+\d+\s*\*\s*(\d+(?:\.\d+)?)')
COST_FLAT = re.compile(r'"COST:\s*"\s*&\s*(\d+(?:\.\d+)?)\s*\)')
PER_LEVEL = re.compile(r"^=\s*[A-Z]+\d+\s*\*\s*(\d+(?:\.\d+)?)\s*$")
IF_TRUE = re.compile(r'IF\(\s*[A-Z]+\d+\s*=\s*TRUE\s*,\s*"([^"]*)"\s*,\s*"([^"]*)"\s*\)', re.IGNORECASE)


def _number(value):
    number = float(value)
    return int(number) if number.is_integer() else number


def _if_texts(value):
    """(text when obtained, text before) from an IF(cell=TRUE, "a", "b") formula."""
    match = IF_TRUE.search(value) if isinstance(value, str) else None
    return (match.group(1), match.group(2)) if match else (None, text(value))


def _strip_bonus(value):
    return re.sub(r"^BONUS:\s*", "", value).strip() if value else None


def _level_limits(sheet):
    """{(row, col): max} from whole-number validations on input cells."""
    limits = {}
    for validation in sheet.data_validations.dataValidation:
        if validation.formula2 is None:
            continue
        try:
            maximum = _number(validation.formula2)
        except ValueError:
            continue
        for block in str(validation.sqref).split():
            min_col, min_row, max_col, max_row = range_boundaries(block)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    limits[(row, col)] = maximum
    return limits


def _fill(sheet, row, col):
    cell = sheet.cell(row, col)
    return str(cell.fill.fgColor.rgb) if cell.fill is not None and cell.fill.fill_type else None


def _pages(sheet):
    pages = []
    for col in range(1, sheet.max_column + 1):
        match = re.fullmatch(r"PAGE\s+(\d+)", text(sheet.cell(4, col).value) or "")
        if match:
            pages.append((int(match.group(1)), col))
    if not pages:
        raise MissingHeader(f"{sheet.title}: no 'PAGE n' headers on row 4")
    return sorted(pages)


def _is_anchor(value):
    return isinstance(value, str) and (
        "COST:" in value or "REQUIRES:" in value or text(value) == "REWARD:"
    )


def _node(sheet, formulas_row, col, first_col, limits, page):
    """One node anchored at (formulas_row, col)."""
    row = formulas_row
    anchor = sheet.cell(row, col).value
    detail = sheet.cell(row + 1, col).value
    player = sheet.cell(row + 2, col).value
    label = text(sheet.cell(row - 1, col).value) or text(sheet.cell(row - 1, col + 1).value)

    node = {
        "id": f"{page}-{get_column_letter(col)}{row}",
        "x": col - first_col,
        "y": row - 1 - FIRST_ROW,
        "label": label,
        "kind": "check",
        "maxLevel": 1,
        "cost": None,
        "bonus": None,
        "requires": None,
        "reward": None,
    }

    levels = COST_LEVELS.search(anchor) if isinstance(anchor, str) else None
    if levels:
        maximum = limits.get((row + 2, col))
        if maximum is None:
            raise ValueError(f"{sheet.title} {get_column_letter(col)}{row + 2}: level node has no max level")
        per_level = PER_LEVEL.match(str(sheet.cell(row + 1, col + 3).value or ""))
        if per_level is None:
            raise ValueError(f"{sheet.title} {get_column_letter(col + 3)}{row + 1}: bonus is not level * amount")
        bonus_cell = sheet.cell(row + 1, col + 3)
        node.update({
            "kind": "level",
            "maxLevel": maximum,
            "cost": {"base": _number(levels.group(1)), "perLevel": -_number(levels.group(2))},
            "bonus": {
                "perLevel": _number(per_level.group(1)),
                "percent": "%" in (bonus_cell.number_format or ""),
            },
        })
        return node

    if text(anchor) == "REWARD:":
        node["reward"] = text(detail)
        return node

    flat = COST_FLAT.search(anchor) if isinstance(anchor, str) else None
    obtained, before = _if_texts(detail)
    if flat:
        node["cost"] = {"base": _number(flat.group(1)), "perLevel": 0}
        node["requires"] = before
        node["bonus"] = {"text": _strip_bonus(obtained)}
    else:
        # IF(obtained, "BONUS: 3x DMG", "REQUIRES:") above IF(obtained, effect, requirement)
        headline, _ = _if_texts(anchor)
        node["requires"] = before
        node["bonus"] = {"text": " · ".join(t for t in (_strip_bonus(headline), obtained) if t)}
    if not isinstance(player, bool) and player is not None:
        raise ValueError(f"{sheet.title} {get_column_letter(col)}{row + 2}: expected a checkbox")
    return node


def _page_bottom(sheet, first_col):
    """Last row that still has the page's black background."""
    bottom = FIRST_ROW
    for row in range(FIRST_ROW, sheet.max_row + 1):
        if any(_fill(sheet, row, col) not in (LINK_FILL, None) for col in range(first_col, first_col + PAGE_WIDTH - 1)):
            bottom = row
    return bottom


def _links(sheet, nodes, first_col, bottom):
    """Connectors: grey cells grouped by adjacency, with the nodes each touches."""
    covered = {}
    for node in nodes:
        for dy in range(NODE_HEIGHT):
            for dx in range(NODE_WIDTH):
                covered[(node["y"] + dy, node["x"] + dx)] = node["id"]

    grey = {
        (row - FIRST_ROW, col - first_col)
        for row in range(FIRST_ROW, bottom + 1)
        for col in range(first_col, first_col + PAGE_WIDTH - 1)
        if _fill(sheet, row, col) == LINK_FILL and (row - FIRST_ROW, col - first_col) not in covered
    }

    links, seen = [], set()
    for start in sorted(grey):
        if start in seen:
            continue
        stack, cells, touching = [start], [], set()
        seen.add(start)
        while stack:
            y, x = stack.pop()
            cells.append([y, x])
            for ny, nx in ((y - 1, x), (y + 1, x), (y, x - 1), (y, x + 1)):
                if (ny, nx) in grey and (ny, nx) not in seen:
                    seen.add((ny, nx))
                    stack.append((ny, nx))
                elif (ny, nx) in covered:
                    touching.add(covered[(ny, nx)])
        if len(touching) >= 2:
            links.append({"cells": sorted(cells), "nodes": sorted(touching)})
    return links


def extract_mastery(sheet):
    """Return (pages, icons); icons maps a file name to image bytes, shared by nodes."""
    limits = _level_limits(sheet)
    images = images_by_cell(sheet)
    pages, icons = [], {}

    for page, first_col in _pages(sheet):
        bottom = _page_bottom(sheet, first_col)
        nodes = []
        for row in range(FIRST_ROW, bottom + 1):
            for col in range(first_col, first_col + PAGE_WIDTH - 1):
                if _is_anchor(sheet.cell(row, col).value):
                    nodes.append(_node(sheet, row, col, first_col, limits, page))

        for node in nodes:
            row, col = node["y"] + FIRST_ROW, node["x"] + first_col
            inside = sorted(
                (r, c) for (r, c) in images
                if row <= r < row + NODE_HEIGHT and col <= c < col + NODE_WIDTH
            )
            node["icon"] = None
            node["badge"] = None
            for position in inside:
                data = images[position]
                name = hashlib.sha1(data).hexdigest()[:12]
                icons[name] = data
                if position == (row, col) and node["icon"] is None:
                    node["icon"] = name
                elif node["badge"] is None:
                    node["badge"] = name
            if node["icon"] is None and node["badge"] is not None:
                node["icon"], node["badge"] = node["badge"], None

        pages.append({
            "page": page,
            "columns": PAGE_WIDTH - 1,
            "rows": bottom - FIRST_ROW + 1,
            "nodes": nodes,
            "links": _links(sheet, nodes, first_col, bottom),
        })
    return pages, icons
