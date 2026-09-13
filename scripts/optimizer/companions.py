"""COMPANIONS, Companions Data and Sprites -> companions, their passive
skills with costs, advancement skins and promotion options.

COMPANIONS lays the four companions side by side, nine columns apart. Each
has nine passive skills in three groups (rows "PASSIVE I/II/III"): a name
cell ("Intensive Fire\\nMax Lv.100", or an IF formula that locks it below an
advancement), a level input, an effect label and an effect formula.

Companions Data holds each companion's cost table (stones and emeralds per
level, skills in the same order) and the promotion tables. Sprites lists
each companion's advancement skins ("Elf 000" ...) with their art.
"""

import re

from openpyxl.utils import get_column_letter

from optimizer.workbook import MissingHeader, find_cell, images_by_cell, number, text

GROUP_TITLES = ["PASSIVE I", "PASSIVE II", "PASSIVE III"]
SKILL_ROWS_PER_GROUP = 3
RANKS = ["1st", "2nd", "3rd", "4th", "5th", "6th", "7th"]
RANK_MULTIPLIERS = {rank: 1 + 0.5 * i for i, rank in enumerate(RANKS)}

NAME_MAX = re.compile(r'"?([^"\n]+)\n\s*Max Lv\.?\s*(\d+)"?')
LOCK = re.compile(r"<\s*(\d+)")
IF_LAST = re.compile(r"^=\s*IF\(.*?,\s*0\s*,\s*(.+)\)\s*$", re.IGNORECASE | re.DOTALL)
CELL_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+")
SAFE = re.compile(r"^[L0-9.+\-*/() ]+$")


def _skill_name(value):
    match = NAME_MAX.search(value or "")
    if not match:
        raise ValueError(f"not a skill name with a max level: {value!r}")
    return match.group(1).strip(), int(match.group(2))


def _unlock(value):
    """(advancement needed, whether every companion needs it) from a lock formula."""
    if not (isinstance(value, str) and value.startswith("=")):
        return None, False
    match = LOCK.search(value)
    return (int(match.group(1)) if match else None), value.upper().lstrip("=").startswith("IF(OR(")


def _effect(formula, number_format):
    """{"kind": "linear", "perLevel": x} or {"kind": "understanding"}, plus how to show it."""
    formula = getattr(formula, "text", formula)  # array formulas keep their source in .text
    display = "flat"
    if '"%"' in number_format:
        display = "percentLiteral"
    elif "%" in number_format:
        display = "percent"

    if isinstance(formula, str) and "SEQUENCE" in formula.upper():
        return {"kind": "understanding", "display": display}

    expression = formula if isinstance(formula, str) else ""
    wrapped = IF_LAST.match(expression)
    expression = (wrapped.group(1) if wrapped else expression.lstrip("=")).strip()
    expression = CELL_REF.sub("L", expression).replace("%", "/100")
    if expression.count("L") != 1 or not SAFE.match(expression):
        raise ValueError(f"effect formula is not level x amount: {formula!r}")

    def at(level):
        return eval(expression, {"__builtins__": {}}, {"L": level})  # noqa: S307 - digits and operators only

    per_level = at(1)
    if abs(at(2) - 2 * per_level) > 1e-9 or at(0) != 0:
        raise ValueError(f"effect formula is not linear in the level: {formula!r}")
    return {"kind": "linear", "perLevel": round(per_level, 10), "display": display}


def _companion_columns(sheet):
    title_row, _ = find_cell(sheet, "PASSIVE I", max_row=40)
    columns = []
    for col in range(1, sheet.max_column + 1):
        if text(sheet.cell(title_row, col).value) == "PASSIVE I":
            columns.append(col)
    if len(columns) != 4:
        raise MissingHeader(f"{sheet.title}: expected 4 'PASSIVE I' headers on row {title_row}, found {len(columns)}")
    return columns


def _skills(sheet, title_col):
    """Nine skills below the PASSIVE I/II/III titles in one companion's block."""
    skills = []
    row = 1
    for group, title in enumerate(GROUP_TITLES, start=1):
        while row <= sheet.max_row and text(sheet.cell(row, title_col).value) != title:
            row += 1
        if row > sheet.max_row:
            raise MissingHeader(f"{sheet.title} column {get_column_letter(title_col)}: no {title!r}")
        first = row + 2  # title, then the PASSIVE ABILITY / LEVEL / EFFECTS header
        for offset in range(SKILL_ROWS_PER_GROUP):
            r = first + offset
            name_cell = sheet.cell(r, title_col + 1).value
            name, max_level = _skill_name(name_cell)
            unlock, all_companions = _unlock(name_cell)
            effect_cell = sheet.cell(r, title_col + 4)
            skills.append({
                "name": name,
                "group": group,
                "maxLevel": max_level,
                "effect": text(sheet.cell(r, title_col + 3).value),
                "formula": _effect(effect_cell.value, effect_cell.number_format or ""),
                "unlock": {"advancement": unlock, "allCompanions": all_companions} if unlock else None,
            })
        row = first + SKILL_ROWS_PER_GROUP
    return skills


def _cost_tables(data_sheet, names):
    """{companion: [[(stones, emeralds) per level] per skill]} from the Level blocks on row 2."""
    level_cols = [c for c in range(1, data_sheet.max_column + 1) if text(data_sheet.cell(2, c).value) == "Level"]
    if len(level_cols) < len(names):
        raise MissingHeader(f"{data_sheet.title}: expected {len(names)} 'Level' cost blocks on row 2")
    tables = {}
    for name, col in zip(names, level_cols):
        skills = [[] for _ in range(9)]
        row = 4
        while isinstance(number(data_sheet.cell(row, col).value), int):
            level = number(data_sheet.cell(row, col).value)
            if level != len(skills[0]):
                raise ValueError(f"{data_sheet.title} row {row}: {name} cost table expected level {len(skills[0])}")
            for i in range(9):
                stones = number(data_sheet.cell(row, col + 1 + 2 * i).value) or 0
                emeralds = number(data_sheet.cell(row, col + 2 + 2 * i).value) or 0
                skills[i].append([stones, emeralds])
            row += 1
        tables[name] = skills
    return tables


def _promotion(data_sheet):
    title_row, title_col = find_cell(data_sheet, "PROMOTION OPTIONS TABLE", max_row=1)
    header_row = title_row + 1
    headers = [text(data_sheet.cell(header_row, c).value) for c in range(title_col, title_col + 7)]
    if headers[:2] != ["ID", "Colour"]:
        raise MissingHeader(f"{data_sheet.title}: promotion table must start with ID, Colour")
    options = [h for h in headers[2:] if h and h != "Locked"]
    tiers = []
    row = header_row + 1
    while isinstance(number(data_sheet.cell(row, title_col).value), int):
        tiers.append({
            "colour": text(data_sheet.cell(row, title_col + 1).value),
            "values": {opt: number(data_sheet.cell(row, title_col + 2 + i).value) for i, opt in enumerate(options)},
        })
        row += 1

    rank_row, rank_col = find_cell(data_sheet, "Promotion #")
    slots = []
    row = rank_row + 1
    while isinstance(number(data_sheet.cell(row, rank_col).value), int):
        ranks = [text(data_sheet.cell(row, rank_col + 1 + i).value) for i in range(7)]
        slots.append([rank if rank in RANKS else None for rank in ranks])
        row += 1
    return {
        "options": options,
        "tiers": tiers,
        "rankMultipliers": RANK_MULTIPLIERS,
        # index 0 is advancement 000; each entry is the 7 slots' ranks (null = locked)
        "slotsByAdvancement": slots,
    }


def _skins(sprites_sheet, name):
    """[(skin name, image bytes or None)] under a companion's name on Sprites row 1."""
    _, col = find_cell(sprites_sheet, name, max_row=1)
    images = images_by_cell(sprites_sheet)
    skins = []
    row = 2
    while (skin := text(sprites_sheet.cell(row, col).value)) is not None:
        skins.append((skin, images.get((row, col + 1))))
        row += 1
    return skins


def extract_companions(sheet, data_sheet, sprites_sheet):
    """Return (companions, promotion, art); art maps a file stem to image bytes."""
    columns = _companion_columns(sheet)
    title_row, _ = find_cell(sheet, "PASSIVE I", max_row=40)
    promotion = _promotion(data_sheet)

    companions, art = [], {}
    names = []
    for col in columns:
        # The name sits just above the "SKIN :" label of the companion's info block.
        skin_row, _ = find_cell(sheet, "SKIN :", min_col=col + 2, max_col=col + 2, max_row=title_row)
        name = text(sheet.cell(skin_row - 1, col + 1).value)
        if not name:
            raise MissingHeader(f"{sheet.title} column {get_column_letter(col + 1)}: no companion name")
        names.append(name)

    costs = _cost_tables(data_sheet, names)
    for index, (col, name) in enumerate(zip(columns, names)):
        element_row, element_col = find_cell(sheet, "ELEMENT :", min_col=col, max_col=col + 3, max_row=title_row)
        skills = _skills(sheet, col)
        for skill, table in zip(skills, costs[name]):
            skill["costs"] = table

        skins = []
        for skin, data in _skins(sprites_sheet, name):
            match = re.search(r"(\d{3})$", skin)
            advancement = int(match.group(1)) if match else len(skins)
            stem = f"{name.lower()}-{advancement:03d}"
            if data is not None:
                art[stem] = data
            skins.append({"advancement": advancement, "name": skin, "icon": stem if data is not None else None})

        companions.append({
            "id": index,
            "name": name,
            "element": text(sheet.cell(element_row, element_col + 1).value),
            "skills": skills,
            "skins": skins,
        })
    return companions, promotion, art
