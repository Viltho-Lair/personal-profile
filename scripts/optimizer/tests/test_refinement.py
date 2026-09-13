import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from optimizer.refinement import extract_refinement


def refinement_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKILLS"
    cells = {
        "B95": "SKILL REFINEMENT",
        "D97": "Fire Slash", "I97": "Flame Slash",
        "D99": "Refinement Effect", "I99": "Refinement Effect",
        "D105": "Owned Effects", "I105": "Owned Effects",
        "D106": "Character HP", "I106": "CRIT Dmg",
    }
    for ref, value in cells.items():
        sheet[ref] = value
    for ref, items in (("F106", '"0%,2%,4%,6%"'), ("K106", '"0,20,40,60"')):
        validation = DataValidation(type="list", formula1=items)
        validation.add(ref)
        sheet.add_data_validation(validation)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)["SKILLS"]


class Refinement(unittest.TestCase):
    def test_lines_by_grade_and_owned_values(self):
        skills = {"Fire Slash": {"grade": "Common"}, "Flame Slash": {"grade": "Rare"}}
        data = extract_refinement(refinement_sheet(), skills)
        self.assertEqual(data["skills"], [
            {"name": "Fire Slash", "lines": 3, "owned": {"stat": "Character HP", "values": [0.0, 0.02, 0.04, 0.06], "percent": True}},
            {"name": "Flame Slash", "lines": 5, "owned": {"stat": "CRIT Dmg", "values": [0.0, 20.0, 40.0, 60.0], "percent": False}},
        ])
        self.assertEqual(data["options"]["DMG Increase(%)"][5], [13, 20])
