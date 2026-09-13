import unittest

from openpyxl.utils import get_column_letter

from optimizer.companions import _effect, extract_companions
from optimizer.tests.support import build_sheet

NAMES = ["Ellie", "Zeke", "Miho", "Luna"]


def companions_sheet(overrides=None):
    cells = {}
    for i, name in enumerate(NAMES):
        col = 3 + 9 * i  # C, L, U, AD
        c = lambda offset, row: f"{get_column_letter(col + offset)}{row}"  # noqa: E731
        cells[c(1, 5)] = name
        cells[c(2, 6)] = "SKIN :"
        cells[c(2, 7)] = "ELEMENT :"
        cells[c(3, 7)] = "Wind"
        for group, title_row in enumerate((14, 20, 26), start=1):
            cells[c(0, title_row)] = f"PASSIVE {'I' * group}"
            for k in range(3):
                row = title_row + 2 + k
                level_ref = f"{get_column_letter(col + 2)}{row}"
                if group == 2 and k == 0:
                    cells[c(1, row)] = f'=IF(OR(INT(RIGHT(F6,3))<13, INT(RIGHT(O6,3))<13), "Locked", "Skill {group}{k}\nMax Lv.100")'
                elif group == 2 and k == 2:
                    cells[c(1, row)] = f'=IF(INT(RIGHT(F6,3))<17, "Locked", "Understanding\nMax Lv.1500")'
                else:
                    cells[c(1, row)] = f"Skill {group}{k}\nMax Lv.100"
                cells[c(3, row)] = "Extra ATK"
                cells[c(4, row)] = (
                    f"=IF(E{row}>10, SUM(SEQUENCE(D43-1,1,10,10)), {level_ref})/100" if group == 2 and k == 2
                    else f"={level_ref}*0.03"
                )
    cells.update(overrides or {})
    return build_sheet(cells, title="COMPANIONS")


def data_sheet():
    cells = {
        "X1": "PROMOTION OPTIONS TABLE", "X2": "ID", "Y2": "Colour", "Z2": "Extra ATK", "AA2": "Extra EXP",
        "AB2": "Monster Gold", "AC2": "Extra HP", "AD2": "Locked",
        "X3": 1, "Y3": "White", "Z3": 0.03, "AA3": 0.01, "AB3": 0.03, "AC3": 0.05, "AD3": 0,
        "X4": 2, "Y4": "Green", "Z4": 0.04, "AA4": 0.02, "AB4": 0.04, "AC4": 0.07, "AD4": 0,
        "X12": "Promotion #", "X13": 1, "Y13": "1st", "X14": 2, "Y14": "2nd", "Z14": "1st",
        "U6": "ELEMENT DMG  INCREMENTS", "U7": 1, "V7": 0.005, "U8": 2, "V8": 0.01,
    }
    for i in range(4):
        col = 33 + 20 * i  # AG, BA, BU, CO
        cells[f"{get_column_letter(col)}2"] = "Level"
        for level in range(3):
            cells[f"{get_column_letter(col)}{4 + level}"] = level
            for skill in range(9):
                cells[f"{get_column_letter(col + 1 + 2 * skill)}{4 + level}"] = level * 10
                cells[f"{get_column_letter(col + 2 + 2 * skill)}{4 + level}"] = level * 5
    return build_sheet(cells, title="Companions Data")


def sprites_sheet():
    cells = {"A1": "Ellie", "C1": "Zeke", "E1": "Miho", "G1": "Luna",
             "A2": "Elf 000", "A3": "Archer 001", "C2": "Tramp Swordsman 000", "E2": "Looter 000",
             "G2": "Trainee Magician 000"}
    return build_sheet(cells, title="Sprites", images=[("B2", 128), ("B3", 128)])


class ExtractCompanions(unittest.TestCase):
    def setUp(self):
        self.companions, self.promotion, self.art = extract_companions(companions_sheet(), data_sheet(), sprites_sheet())

    def test_four_companions_with_nine_skills_each(self):
        self.assertEqual([c["name"] for c in self.companions], NAMES)
        self.assertTrue(all(len(c["skills"]) == 9 for c in self.companions))
        self.assertEqual(self.companions[0]["element"], "Wind")

    def test_skill_name_max_level_unlock_and_effect(self):
        skills = self.companions[0]["skills"]
        self.assertEqual((skills[0]["name"], skills[0]["maxLevel"], skills[0]["unlock"]), ("Skill 10", 100, None))
        self.assertEqual(skills[3]["unlock"], {"advancement": 13, "allCompanions": True})
        self.assertEqual((skills[5]["name"], skills[5]["maxLevel"]), ("Understanding", 1500))
        self.assertEqual(skills[5]["unlock"], {"advancement": 17, "allCompanions": False})
        self.assertEqual(skills[0]["formula"]["perLevel"], 0.03)
        self.assertEqual(skills[5]["formula"]["kind"], "understanding")

    def test_costs_per_level_in_skill_order(self):
        self.assertEqual(self.companions[1]["skills"][8]["costs"], [[0, 0], [10, 5], [20, 10]])

    def test_skins_by_advancement_with_art(self):
        skins = self.companions[0]["skins"]
        self.assertEqual([(s["advancement"], s["name"], s["icon"]) for s in skins],
                         [(0, "Elf 000", "ellie-000"), (1, "Archer 001", "ellie-001")])
        self.assertIsNone(self.companions[1]["skins"][0]["icon"])
        self.assertEqual(sorted(self.art), ["ellie-000", "ellie-001"])

    def test_promotion_tiers_and_slot_ranks(self):
        self.assertEqual(self.promotion["options"][:3], ["Extra ATK", "CRIT Dmg", "Extra HP"])
        self.assertEqual(len(self.promotion["options"]), 11)
        green = self.promotion["tiers"][1]
        self.assertEqual(green["colour"], "Green")
        self.assertEqual({k: green["values"][k] for k in ["Extra ATK", "Extra EXP", "Monster Gold", "Extra HP"]},
                         {"Extra ATK": 0.04, "Extra EXP": 0.02, "Monster Gold": 0.04, "Extra HP": 0.07})
        self.assertEqual(green["values"]["Accuracy"], 4)
        self.assertEqual(self.promotion["flatOptions"], ["Accuracy", "Dodge", "CC Resist"])
        self.assertEqual(self.promotion["slotsByAdvancement"][1][:3], ["2nd", "1st", None])
        self.assertEqual(self.promotion["rankMultipliers"]["7th"], 4)
        self.assertEqual(self.promotion["elementIncrements"], [0.005, 0.01])


class Effect(unittest.TestCase):
    def test_formats_and_linear_amounts(self):
        self.assertEqual(_effect("=N17*0.04*100", '"+"0"%"'), {"kind": "linear", "perLevel": 4.0, "display": "percentLiteral"})
        self.assertEqual(_effect('=if(D22="Locked", 0, E22/100)', '"+"0%')["perLevel"], 0.01)
        self.assertEqual(_effect("=E28*1%", "0%")["perLevel"], 0.01)
        self.assertEqual(_effect("=W17*3", "+0")["display"], "flat")

    def test_rejects_non_linear_formulas(self):
        with self.assertRaises(ValueError):
            _effect("=E16*E16", "0")


if __name__ == "__main__":
    unittest.main()
