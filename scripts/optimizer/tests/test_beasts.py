import unittest

from openpyxl.utils import get_column_letter

from optimizer.beasts import extract_beasts
from optimizer.tests.support import build_sheet


class Beasts(unittest.TestCase):
    def test_beast_rows_and_affection_table(self):
        cells = {
            "A2": "Tier", "B2": "Full Name", "C2": "Skill Description", "D2": "Skill X", "E2": "None",
            "M2": "1st Effect Name", "N2": "2nd Effect Name",
            "S2": "Type Icon", "T2": "Icon 1", "U2": "Icon 2", "V2": "Icon None",
            "A4": "Common", "B4": "Gray Wolf", "C4": "After X strike skills used, ATK +Y% for 10s", "D4": 9, "E4": 0,
            "M4": "Increased Attack :",
            "A5": "Unique", "B5": "Light Draco", "C5": "When maxing stacks of X, Boss DMG increases by Y% for 60 seconds",
            "D5": "all buffs", "E5": 0, "M5": "Increased MSPD :", "N5": "Increased Affection :",
        }
        for a in range(7):
            cells[f"{get_column_letter(6 + a)}4"] = 10 + a
            cells[f"{get_column_letter(6 + a)}5"] = 30 + a
        # Stat table: level column P, Common block from Q, Unique block right after it.
        cells["Q11"] = "Common"
        cells[f"{get_column_letter(17 + 28)}11"] = "Unique"
        for block in range(2):
            for a in range(7):
                for s, name in enumerate(["Combat", "Draco Combat", "ATK/Affection", "MSPD"]):
                    col = get_column_letter(17 + block * 28 + a * 4 + s)
                    cells[f"{col}15"] = name
                    for level in range(1, 10 * (a + 1) + 1 if a < 1 else 3):
                        cells[f"{col}{15 + level}"] = block * 100 + a * 10 + s + level / 100
        for level in range(1, 11):
            cells[f"P{15 + level}"] = level
        sheet = build_sheet(cells, title="Companions Data", images=[("T4", 32), ("V4", 16), ("G3", 16)])
        beasts, art = extract_beasts(sheet)
        wolf, draco = beasts["beasts"]
        self.assertEqual(wolf["family"], "Wolf")
        self.assertEqual(wolf["mounted"], ["atk"])
        self.assertEqual(wolf["skill"]["values"], [10, 11, 12, 13, 14, 15, 16])
        self.assertEqual(draco["mounted"], ["mspd", "affection"])
        self.assertEqual(draco["skill"]["x"], "all buffs")
        self.assertEqual(len(beasts["tables"]["Common"]["combat"][0]), 10)
        self.assertEqual(beasts["tables"]["Unique"]["mspd"][1][:2], [113.01, 113.02])
        self.assertEqual(wolf["art"], {"sprite": "gray-wolf-sprite", "egg": "gray-wolf-egg"})
        self.assertEqual(draco["art"], {})
        self.assertEqual(beasts["awakenIcons"][:2], ["awaken-1", None])
        self.assertEqual(sorted(art), ["awaken-1", "gray-wolf-egg", "gray-wolf-sprite"])
